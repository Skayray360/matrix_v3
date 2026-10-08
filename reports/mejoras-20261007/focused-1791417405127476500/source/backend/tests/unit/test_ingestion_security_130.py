# Creado por Aldo Garcia.
"""Regresiones de ingesta 1.3.0: archivos sinteticos, sin servicios externos."""

from __future__ import annotations

import io
import zipfile
from dataclasses import replace

import pytest

from app.common.errors import UnsupportedFileError
from app.ingestion import isolated_extraction, loaders
from app.rag.schemas import SCOPE_CONVERSATION
from app.rag.vector_store import payload_to_evidence
from app.security.upload_guard import validate_upload
from tests.unit.test_rag_pipeline_isolated import make_chunk

pytestmark = pytest.mark.unit


@pytest.mark.parametrize("filename", ["politica.pdf", "politica.docx", "politica.xlsx"])
def test_corporate_extraction_rejects_spoofed_type_before_starting_parser(monkeypatch, filename):
    monkeypatch.setattr(
        isolated_extraction.subprocess, "Popen",
        lambda *args, **kwargs: pytest.fail("El contenido invalido llego al parser."),
    )
    with pytest.raises(UnsupportedFileError):
        isolated_extraction.extract_document(b"MZ contenido sintetico", filename=filename)


def test_corporate_extraction_rejects_ooxml_zip_bomb_before_starting_parser(monkeypatch):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "X" * 1_000_000)
    monkeypatch.setattr(
        isolated_extraction.subprocess, "Popen",
        lambda *args, **kwargs: pytest.fail("La bomba comprimida llego al parser."),
    )
    with pytest.raises(UnsupportedFileError, match="compresion"):
        isolated_extraction.extract_document(buffer.getvalue(), filename="politica.docx")


def test_long_windows_filename_keeps_its_validated_extension():
    result = validate_upload(b"%PDF-1.7\ncontenido sintetico", filename="a" * 240 + ".pdf")
    assert result.extension == ".pdf"
    assert len(result.safe_display_name) == 200
    assert result.safe_display_name.endswith(".pdf")


def test_emptied_corporate_text_keeps_the_existing_retirement_flow():
    # El servicio debe recibir "empty" para desactivar evidencia anterior, no
    # una excepcion que haga rollback y conserve la version vaciada en disco.
    result = isolated_extraction.extract_document(b"", filename="politica.md")
    assert result.is_empty


def test_sparse_xlsx_does_not_read_cells_beyond_physical_row_limit(monkeypatch):
    from openpyxl import Workbook

    monkeypatch.setattr(loaders, "MAX_XLSX_ROWS_PER_SHEET", 3)
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(1, 1, "puesto")
    sheet.cell(2, 1, "permitido")
    sheet.cell(5, 1, "fuera_del_limite")
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = loaders.extract_xlsx(buffer.getvalue())
    assert "permitido" in result.plain_text()
    assert "fuera_del_limite" not in result.plain_text()
    assert any("filas" in warning for warning in result.warnings)


def test_xlsx_reports_columns_omitted_by_the_operational_limit(monkeypatch):
    from openpyxl import Workbook

    monkeypatch.setattr(loaders, "MAX_XLSX_COLUMNS", 2)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["puesto", "area", "dato_no_leido"])
    sheet.append(["analista", "sistemas", "fuera_del_limite"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = loaders.extract_xlsx(buffer.getvalue())
    assert "analista" in result.plain_text()
    assert "fuera_del_limite" not in result.plain_text()
    assert any("columnas" in warning for warning in result.warnings)


def test_small_xlsx_does_not_report_false_truncation():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["puesto"])
    sheet.append(["analista"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = loaders.extract_xlsx(buffer.getvalue())
    assert "puesto: analista" in result.plain_text()
    assert result.warnings == []


def test_corporate_citations_distinguish_equal_filenames_in_different_folders():
    original = make_chunk("Regla A", category="prestaciones", document_id="doc-a")
    first = replace(original.metadata, filename="politica.pdf", relative_path="prestaciones/a/politica.pdf")
    second = replace(first, document_id="doc-b", relative_path="prestaciones/b/politica.pdf")
    first_result = payload_to_evidence(first.to_payload("Regla A"), 0.9)
    second_result = payload_to_evidence(second.to_payload("Regla B"), 0.9)

    assert first_result.source_id == first.source_id
    assert second_result.source_id == second.source_id
    assert first_result.source_id != second_result.source_id


def test_private_citations_distinguish_two_uploads_with_the_same_display_name():
    original = make_chunk(
        "CV A", category="__private__", document_id="doc-a", scope=SCOPE_CONVERSATION,
        owner="user-a", conversation="conversation-a",
    )
    first = replace(original.metadata, filename="cv.pdf", relative_path="cv.pdf")
    second = replace(first, document_id="doc-b")
    first_result = payload_to_evidence(first.to_payload("CV A"), 0.9)
    second_result = payload_to_evidence(second.to_payload("CV B"), 0.9)

    assert first_result.source_id == first.source_id
    assert second_result.source_id == second.source_id
    assert first_result.source_id != second_result.source_id


def test_common_corporate_citation_preserves_its_existing_id():
    chunk = make_chunk("Regla A", category="prestaciones")
    assert chunk.metadata.source_id == "prestaciones/prestaciones.md#0"
    assert payload_to_evidence(chunk.metadata.to_payload(chunk.text), 0.9).source_id == chunk.metadata.source_id
