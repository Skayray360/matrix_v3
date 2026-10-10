# Creado por Aldo Garcia.
"""Resumenes por identidad documental con SQL y Qdrant locales sinteticos."""

from __future__ import annotations

import io
from contextlib import contextmanager
from dataclasses import replace
from types import SimpleNamespace

import docx
import pytest
from qdrant_client import QdrantClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.agents.orchestrator import Orchestrator
from app.common.ids import new_id, utcnow_naive
from app.config.settings import Settings
from app.database.models import Base, Document
from app.ingestion.loaders import extract_docx
from app.llm.model_policy import Intent
from app.rag.chunking import chunk_blocks
from app.rag.document_selection import explicit_document_filenames, has_document_reference
from app.rag.index_manifest import indexing_fingerprint
from app.rag.retriever import Retriever
from app.rag.schemas import Chunk, ChunkMetadata
from app.rag.vector_store import VectorStore
from tests.conftest import make_context

pytestmark = pytest.mark.unit


class NoEmbedding:
    def embed_one(self, *_args, **_kwargs):
        raise AssertionError("Un documento identificado no requiere busqueda semantica")


@pytest.fixture
def corpus(monkeypatch):
    settings = Settings(_env_file=None, app_env="test", rag_embedding_dimension=3, ollama_embedding_dimension=3)
    for module in ("retriever", "vector_store", "index_manifest"):
        monkeypatch.setattr(f"app.rag.{module}.get_settings", lambda: settings)
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    factory = sessionmaker(engine, expire_on_commit=False)

    @contextmanager
    def scope():
        with factory() as db:
            yield db
            db.commit()

    monkeypatch.setattr("app.rag.index_manifest.session_scope", scope)
    store = VectorStore(client=QdrantClient(location=":memory:"))
    fingerprint = indexing_fingerprint()

    def publish(filename, *, count=3, scope_name="conversation", user="owner", conversation="thread",
                category="tecnologia", generation=None, fingerprint_value=None, blocks=None):
        document_id = new_id()
        generation = generation or new_id()
        active_fingerprint = fingerprint_value or fingerprint
        texts = [f"Paso sintetico {index}: tarea de {filename}." for index in range(count)]
        if blocks is not None:
            texts = [block.text for block in blocks]
        chunks = [Chunk(text, ChunkMetadata(
            chunk_id=new_id(), document_id=document_id, document_sha256="a" * 64,
            filename=filename, relative_path=f"{category}/{filename}", category=category, subpath="",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            page_or_sheet=blocks[index].page_or_sheet if blocks is not None else f"parrafo {index + 1}",
            section=blocks[index].section if blocks is not None else f"Seccion {index // 10}",
            chunk_index=index, embedding_model=settings.ollama_embedding_model,
            embedding_dimension=3, ingestion_version="synthetic", generation=generation,
            index_fingerprint=active_fingerprint, scope=scope_name,
            owner_user_id=user if scope_name == "conversation" else None,
            conversation_id=conversation if scope_name == "conversation" else None,
        )) for index, text in enumerate(texts)]
        store.upsert_chunks(chunks, [[1.0, 0.0, 0.0] for _ in chunks])
        with scope() as db:
            db.add(Document(
                id=document_id, scope=scope_name, category=category, filename=filename,
                mime_type=chunks[0].metadata.mime_type, sha256="a" * 64, status="indexed",
                chunk_count=len(chunks), active_generation=generation, index_fingerprint=active_fingerprint,
                owner_user_id=user if scope_name == "conversation" else None,
                conversation_id=conversation if scope_name == "conversation" else None,
            ))
        return chunks

    yield SimpleNamespace(store=store, settings=settings, publish=publish, scope=scope, fingerprint=fingerprint,
                          retriever=Retriever(store=store, llm=NoEmbedding()), ctx=make_context(user_id="owner"))
    store.close()
    engine.dispose()


@pytest.mark.parametrize("question", [
    "explicame sobre este documento Guia_configuracion_Kerberos_AD_DOMINIO_MX.",
    "Resume Guía configuración Kerberos AD DOMINIO MX",
    "Dame un resumen de Guia_configuracion_Kerberos_AD_DOMINIO_MX.docx",
])
def test_complete_private_docx_has_85_ordered_citable_paragraphs_without_embedding(corpus, question):
    document = docx.Document()
    for index in range(85):
        document.add_paragraph(f"Paso sintetico {index}: configurar componente autorizado.")
    buffer = io.BytesIO()
    document.save(buffer)
    blocks = chunk_blocks(extract_docx(buffer.getvalue()).blocks, chunk_size_tokens=900, overlap_tokens=120)
    requested = corpus.publish("Guia_configuracion_Kerberos_AD_DOMINIO_MX.docx", blocks=blocks)
    corpus.publish("Manual_No_Relacionado.docx")
    result = corpus.retriever.retrieve_attachment_summary(ctx=corpus.ctx, conversation_id="thread", question=question)
    assert len(result.evidences) == len(blocks) == 85
    assert [item.source_id for item in result.evidences] == [chunk.metadata.source_id for chunk in requested]
    assert [item.page_or_sheet for item in result.evidences] == [f"parrafo {number}" for number in range(1, 86)]
    assert result.used_private_scope and not result.truncated and not result.clarification


def test_named_corporate_document_bypasses_top_k_and_similarity(corpus):
    corpus.settings.rag_top_k = 1
    corpus.settings.rag_min_similarity = 1.0
    chunks = corpus.publish("Guia_Configuracion_Directorio.docx", count=85, scope_name="corporate")
    result = corpus.retriever.retrieve(
        ctx=corpus.ctx, question="Explicame Guia Configuracion Directorio", summary=True,
        authorized_categories=frozenset({"tecnologia"}), include_private=False,
    )
    assert [e.source_id for e in result.evidences] == [chunk.metadata.source_id for chunk in chunks]
    assert result.authorized_categories == ("tecnologia",) and not result.truncated


def test_generic_single_attachment_is_complete_without_embedding(corpus):
    chunks = corpus.publish("unico.docx", count=12)
    result = corpus.retriever.retrieve_attachment_summary(
        ctx=corpus.ctx, conversation_id="thread", question="dame un resumen",
    )
    assert [e.source_id for e in result.evidences] == [chunk.metadata.source_id for chunk in chunks]


def test_multiple_private_documents_need_selection_unless_plural_requested(corpus):
    corpus.publish("Primero.docx")
    corpus.publish("Segundo.docx")
    ambiguous = corpus.retriever.retrieve_attachment_summary(
        ctx=corpus.ctx, conversation_id="thread", question="explicame este documento",
    )
    assert ambiguous.clarification and not ambiguous.has_evidence
    aggregate = corpus.retriever.retrieve_attachment_summary(
        ctx=corpus.ctx, conversation_id="thread", question="resume todos los adjuntos",
    )
    assert {e.filename for e in aggregate.evidences} == {"Primero.docx", "Segundo.docx"}


@pytest.mark.parametrize("question", [
    "Resume No_Disponible.docx", "Explica la Guia_Confidencial",
    "Resume este documento Guía de instalación inexistente",
])
def test_unknown_named_target_never_falls_back_to_other_private_or_corporate_files(corpus, question):
    corpus.publish("Permitido.docx")
    corpus.publish("Permitido.docx", scope_name="corporate")
    private = corpus.retriever.retrieve_attachment_summary(
        ctx=corpus.ctx, conversation_id="thread", question=question,
    )
    corporate = corpus.retriever.retrieve(
        ctx=corpus.ctx, question=question, summary=True, include_private=False,
        authorized_categories=frozenset({"tecnologia"}),
    )
    assert not private.has_evidence and not corporate.has_evidence


@pytest.mark.parametrize("scope", ["conversation", "corporate"])
def test_same_filename_in_multiple_documents_requires_version_selection(corpus, scope):
    corpus.publish("Guia Configuracion.docx", scope_name=scope)
    corpus.publish("Guia Configuracion.docx", scope_name=scope)
    if scope == "conversation":
        result = corpus.retriever.retrieve_attachment_summary(
            ctx=corpus.ctx, conversation_id="thread", question="resume Guia Configuracion.docx",
        )
    else:
        result = corpus.retriever.retrieve(
            ctx=corpus.ctx, question="resume Guia Configuracion.docx", summary=True, include_private=False,
            authorized_categories=frozenset({"tecnologia"}),
        )
    assert result.clarification and not result.has_evidence


def test_metadata_resolution_and_qdrant_both_preserve_owner_conversation_and_category(corpus):
    allowed = corpus.publish("Guia de Configuracion.docx")
    corpus.publish("Guia de Configuracion.docx", user="other-owner")
    corpus.publish("Guia de Configuracion.docx", conversation="other-thread")
    corporate = corpus.publish("Guia de Configuracion.docx", scope_name="corporate")
    corpus.publish("Guia de Configuracion.docx", scope_name="corporate", category="restricted")
    private_docs = corpus.store.list_authorized_documents(
        scope="conversation", user_id="owner", conversation_id="thread", index_fingerprint=corpus.fingerprint,
    )
    assert {doc.document_id for doc in private_docs} == {allowed[0].metadata.document_id}
    private = corpus.retriever.retrieve_attachment_summary(
        ctx=corpus.ctx, conversation_id="thread", question="resume Guia de Configuracion.docx",
    )
    public = corpus.retriever.retrieve(
        ctx=corpus.ctx, question="resume Guia de Configuracion.docx", summary=True, include_private=False,
        authorized_categories=frozenset({"tecnologia"}),
    )
    assert {e.document_id for e in private.evidences} == {allowed[0].metadata.document_id}
    assert {e.document_id for e in public.evidences} == {corporate[0].metadata.document_id}
    assert not corpus.store.list_authorized_documents(scope="corporate", index_fingerprint=corpus.fingerprint)
    assert not corpus.store.list_authorized_documents(
        scope="conversation", user_id="owner", index_fingerprint=corpus.fingerprint,
    )
    # A stale/foreign metadata identity passed by a caller cannot broaden ACL.
    leaked, _ = corpus.store.list_document_chunks(
        documents=private_docs, scope="conversation", user_id="other-owner", conversation_id="thread",
        index_fingerprint=corpus.fingerprint, limit=10,
    )
    assert leaked == []


def test_pending_old_deleted_and_wrong_embedding_generations_never_enter_summary(corpus):
    active = corpus.publish("Guia de Configuracion.docx", count=10)
    unpublished = [replace(chunk, text="NO PUBLICADO", metadata=replace(
        chunk.metadata, chunk_id=new_id(), generation="unpublished-generation",
    )) for chunk in active]
    corpus.store.upsert_chunks(unpublished, [[1.0, 0.0, 0.0] for _ in unpublished])
    corpus.publish("Guia anterior.docx", fingerprint_value="previous-embedding-revision")
    deleted = corpus.publish("Guia de Configuracion.docx")
    pending = corpus.publish("Guia pendiente.docx")
    with corpus.scope() as db:
        db.get(Document, deleted[0].metadata.document_id).deleted_at = utcnow_naive()
        db.get(Document, pending[0].metadata.document_id).status = "pending"
    result = corpus.retriever.retrieve_attachment_summary(
        ctx=corpus.ctx, conversation_id="thread", question="resume Guia de Configuracion.docx",
    )
    assert [e.chunk_id for e in result.evidences] == [chunk.metadata.chunk_id for chunk in active]
    assert not result.truncated


@pytest.mark.parametrize("unavailable", ["fingerprint", "pending"])
def test_unavailable_own_attachment_never_becomes_a_summary_of_another_document(corpus, unavailable):
    private = corpus.publish("Guia de Configuracion.docx")
    corpus.publish("Guia distinta.docx", scope_name="corporate")
    with corpus.scope() as db:
        document = db.get(Document, private[0].metadata.document_id)
        if unavailable == "fingerprint":
            document.index_fingerprint = "previous-revision"
        else:
            document.status = "pending"
            document.active_generation = None
    for question in ("dame un resumen", "explicame Guia de Configuracion.docx"):
        result = corpus.retriever.retrieve(
            ctx=corpus.ctx, conversation_id="thread", question=question,
            authorized_categories=frozenset({"tecnologia"}), summary=True,
        )
        assert not result.has_evidence and "aun no esta disponible" in result.clarification


def test_active_and_unavailable_private_documents_still_require_target_selection(corpus):
    corpus.publish("Guia activa.docx")
    corpus.publish("Guia anterior.docx", fingerprint_value="previous-revision")
    result = corpus.retriever.retrieve_attachment_summary(
        ctx=corpus.ctx, conversation_id="thread", question="dame un resumen",
    )
    assert result.clarification.startswith("Hay varios") and not result.has_evidence


def test_scan_budget_samples_beginning_middle_end_with_explicit_truncation(corpus):
    corpus.settings.rag_summary_scan_max_chunks = 7
    chunks = corpus.publish("Guia Completa.docx", count=85)
    result = corpus.retriever.retrieve_attachment_summary(
        ctx=corpus.ctx, conversation_id="thread", question="resume Guia Completa.docx",
    )
    indices = [int(e.source_id.rsplit("#", 1)[1]) for e in result.evidences]
    assert indices == [0, 14, 28, 42, 56, 70, 84]
    assert result.truncated and result.fetched == 7
    assert result.evidences[0].source_id == chunks[0].metadata.source_id
    assert result.evidences[-1].source_id == chunks[-1].metadata.source_id


def test_different_document_sizes_do_not_force_sampling_when_all_fit(corpus):
    corpus.settings.rag_summary_scan_max_chunks = 100
    corpus.publish("A_grande.docx", count=85)
    corpus.publish("B_pequeno.docx", count=5)
    result = corpus.retriever.retrieve_attachment_summary(
        ctx=corpus.ctx, conversation_id="thread", question="resume todos los documentos",
    )
    assert len(result.evidences) == 90 and not result.truncated


def test_small_second_document_keeps_budget_when_first_document_is_large(corpus):
    corpus.settings.rag_summary_scan_max_chunks = 7
    corpus.publish("A_grande.docx", count=85)
    corpus.publish("B_pequeno.docx", count=1)
    result = corpus.retriever.retrieve_attachment_summary(
        ctx=corpus.ctx, conversation_id="thread", question="resume todos los documentos",
    )
    assert len(result.evidences) == 7 and result.truncated
    assert {e.filename for e in result.evidences} == {"A_grande.docx", "B_pequeno.docx"}


@pytest.mark.parametrize("question", ["dame un resumen", "resume esto", "explicame este documento"])
def test_generic_summary_without_attachment_or_resolved_history_clarifies_before_semantic_search(corpus, question):
    corpus.publish("Documento no seleccionado.docx", scope_name="corporate")
    result = corpus.retriever.retrieve(
        ctx=corpus.ctx, question=question, summary=True, conversation_id="thread",
        authorized_categories=frozenset({"tecnologia"}),
    )
    assert result.clarification and not result.has_evidence


def test_all_corporate_documents_is_not_limited_to_semantic_top_k(corpus):
    corpus.settings.rag_top_k = 1
    corpus.publish("Guia Uno.docx", scope_name="corporate")
    corpus.publish("Guia Dos.docx", scope_name="corporate")
    result = corpus.retriever.retrieve(
        ctx=corpus.ctx, question="resume todos los documentos corporativos", summary=True, include_private=False,
        authorized_categories=frozenset({"tecnologia"}),
    )
    assert len(result.evidences) == 6 and not result.truncated
    assert {e.filename for e in result.evidences} == {"Guia Uno.docx", "Guia Dos.docx"}


def test_mixed_named_private_and_corporate_topic_preserves_both_document_contexts(corpus):
    corpus.settings.rag_top_k = 1
    private = corpus.publish("Guia de Altas y Bajas.docx", count=8)
    public = corpus.publish("Manual General.docx", count=10, scope_name="corporate")
    questions = []

    def embed_one(question):
        questions.append(question)
        return [1.0, 0.0, 0.0]

    orchestrator = Orchestrator.__new__(Orchestrator)
    orchestrator._retriever = Retriever(store=corpus.store, llm=SimpleNamespace(embed_one=embed_one))
    result, _, tools = orchestrator._gather_evidence(
        ctx=corpus.ctx, conversation=SimpleNamespace(id="thread"),
        question="Resume el archivo adjunto Guia de Altas y Bajas.docx y la politica corporativa de vacaciones",
        intent=Intent.DOCUMENT_SUMMARY, authorized_categories=frozenset({"tecnologia"}), comparative=False,
    )
    assert len(result.evidences) == len(private) + len(public)
    assert set(tools) == {"rag", "private_attachment_summary"}
    assert len(questions) == 1 and "Guia de Altas" not in questions[0]
    assert result.used_private_scope and not result.truncated


def test_incomplete_active_vector_generation_reports_coverage_limit(corpus):
    chunks = corpus.publish("Guia Completa.docx", count=5)
    corpus.store._client.delete(
        collection_name=corpus.store.collection_for("conversation"), points_selector=[chunks[-1].metadata.chunk_id],
    )
    result = corpus.retriever.retrieve_attachment_summary(
        ctx=corpus.ctx, conversation_id="thread", question="resume Guia Completa.docx",
    )
    assert len(result.evidences) == 4 and result.truncated


def test_title_match_and_unknown_reference_detector_do_not_use_personal_dates():
    name = "Guía_configuración_AD_Diciembre2022.docx"
    assert explicit_document_filenames("Explica guia configuracion ad diciembre 2022", [name]) == {name}
    assert not explicit_document_filenames("Ingrese en diciembre de 2022", [name])
    assert not has_document_reference("Ingrese en diciembre de 2022")
