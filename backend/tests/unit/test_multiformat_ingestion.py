# Creado por Aldo Garcia.
"""Contratos de extraccion: orden, cifras literales, ubicaciones y fallos acotados."""

from __future__ import annotations

import io
import itertools
import zipfile
from xml.sax.saxutils import escape

import docx
import pytest
from pypdf import PdfWriter

from app.common.errors import ExtractionFailedError, UnsupportedFileError
from app.ingestion import isolated_extraction, loaders, local_ocr, pptx
from app.ingestion.local_ocr import OcrOptions

pytestmark = pytest.mark.unit

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _shape(text: str, *, title: bool = False) -> str:
    placeholder = '<p:ph type="title"/>' if title else ""
    return (f"<p:sp><p:nvSpPr><p:nvPr>{placeholder}</p:nvPr></p:nvSpPr>"
            f"<p:txBody><a:p><a:r><a:t>{escape(text)}</a:t></a:r></a:p></p:txBody></p:sp>")


def presentation(order=(2, 10, 1), *, external=False, dtd=False, empty=False) -> bytes:
    """Paquete sintetico: no usa PowerPoint ni otra libreria como oracle."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        slides = "".join(f'<p:sldId id="{256 + index}" r:id="rel{number}"/>' for index, number in enumerate(order))
        archive.writestr("ppt/presentation.xml",
                         f'<p:presentation xmlns:p="{P}" xmlns:r="{R}"><p:sldIdLst>{slides}</p:sldIdLst></p:presentation>')
        relationships = "".join(
            f'<Relationship Id="rel{number}" Type="{R}/slide" Target="slides/slide{number}.xml"'
            + (' TargetMode="External"' if external else "") + '/>' for number in order
        )
        archive.writestr("ppt/_rels/presentation.xml.rels",
                         f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{relationships}</Relationships>')
        archive.writestr("[Content_Types].xml",
                         '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                         '<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/></Types>')
        for number in order:
            table = '<p:graphicFrame><a:graphic><a:graphicData><a:tbl>'
            for row in (("Concepto", "Monto"), ("A", "1,000,000"), ("B", "1.000.000")):
                table += '<a:tr>' + ''.join(
                    f'<a:tc><a:txBody><a:p><a:r><a:t>{value}</a:t></a:r></a:p></a:txBody></a:tc>' for value in row
                ) + '</a:tr>'
            table += '</a:tbl></a:graphicData></a:graphic></p:graphicFrame>'
            content = "" if empty else _shape(f"Titulo {number}", title=True) + _shape(f"Regla {number}") + table
            declaration = '<!DOCTYPE p:sld [<!ENTITY forbidden "valor">]>' if dtd else ""
            archive.writestr(f"ppt/slides/slide{number}.xml",
                             f'{declaration}<p:sld xmlns:p="{P}" xmlns:a="{A}"><p:cSld><p:spTree>{content}</p:spTree></p:cSld></p:sld>')
    return buffer.getvalue()


@pytest.mark.parametrize("order", list(itertools.permutations((2, 10, 1))))
def test_pptx_preserves_logical_slide_order_and_literal_table_values(order):
    result = loaders.extract_document(presentation(order), filename="sintetica.pptx")
    titles = [block for block in result.blocks if block.kind == "heading"]
    assert [block.text for block in titles] == [f"## Titulo {number}" for number in order]
    assert [block.page_or_sheet for block in titles] == ["diapositiva 1", "diapositiva 2", "diapositiva 3"]
    tables = [block for block in result.blocks if block.kind == "table"]
    assert len(tables) == 3
    assert all(block.text == "| Concepto | Monto |\n| A | 1,000,000 |\n| B | 1.000.000 |" for block in tables)
    assert not result.warnings


@pytest.mark.parametrize("kwargs", [{"external": True}, {"dtd": True}])
def test_pptx_untrusted_relations_and_entities_fail_closed(kwargs):
    with pytest.raises(ExtractionFailedError):
        loaders.extract_pptx(presentation(**kwargs))


def test_pptx_corrupt_zip_is_controlled():
    with pytest.raises(ExtractionFailedError, match="corrupta"):
        loaders.extract_pptx(b"PK\x03\x04 roto")


def test_pptx_image_only_slide_is_not_invented():
    result = loaders.extract_pptx(presentation((1,), empty=True))
    assert result.is_empty
    assert "Diapositiva 1" in result.warnings[0]


def test_pptx_exceeding_slide_limit_is_not_partially_indexed(monkeypatch):
    monkeypatch.setattr(pptx, "MAX_SLIDES", 2)
    with pytest.raises(ExtractionFailedError, match="diapositivas"):
        loaders.extract_pptx(presentation())


def test_pptx_table_limit_is_enforced(monkeypatch):
    monkeypatch.setattr(pptx, "MAX_TABLE_ROWS", 2)
    with pytest.raises(ExtractionFailedError, match="filas"):
        loaders.extract_pptx(presentation((1,)))


@pytest.mark.parametrize("extension", ["doc", "ppt"])
def test_legacy_office_requires_explicit_local_conversion(extension):
    with pytest.raises(UnsupportedFileError, match="conversion local"):
        loaders.extract_document(b"legacy", filename=f"antiguo.{extension}")


@pytest.mark.parametrize("text", ["Nómina y antigüedad", "Monto 1,000,000; tasa 7.5%", "Regla 你好"])
def test_text_bom_encodings_preserve_unicode_and_numbers(text):
    # Propiedad acotada de conservacion frente al cambio de encoding.
    for encoding in ("utf-8-sig", "utf-16", "utf-16-be"):
        data = text.encode(encoding)
        if encoding == "utf-16-be":
            data = b"\xfe\xff" + data
        assert loaders.extract_txt(data).plain_text() == text


def test_docx_locators_are_real_paragraph_and_table_positions():
    document = docx.Document()
    document.add_heading("Reglas", level=1)
    document.add_paragraph("")
    document.add_paragraph("Regla sintetica.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Concepto"
    table.cell(0, 1).text = "Importe"
    document.add_paragraph("Despues de la tabla.")
    buffer = io.BytesIO()
    document.save(buffer)
    result = loaders.extract_docx(buffer.getvalue())
    assert [block.page_or_sheet for block in result.blocks] == ["parrafo 1", "parrafo 3", "tabla 1", "parrafo 4"]
    assert all(block.section == "Reglas" for block in result.blocks)
    assert not any("pagina" in block.page_or_sheet for block in result.blocks)


def test_docx_paragraph_limit_is_explicit(monkeypatch):
    document = docx.Document()
    document.add_paragraph("Permitido")
    document.add_paragraph("Exceso")
    buffer = io.BytesIO()
    document.save(buffer)
    monkeypatch.setattr(loaders, "MAX_DOCX_PARAGRAPHS", 1)
    with pytest.raises(ExtractionFailedError, match="parrafos"):
        loaders.extract_docx(buffer.getvalue())


def _blank_pdf(*, encrypted: bool = False) -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    if encrypted:
        writer.encrypt("sintetico")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_pdf_ocr_is_optional_and_keeps_physical_page_locator(monkeypatch):
    called = []

    def local_page(data, number, options):
        called.append((number, options.language))
        return "Prueba OCR 42"

    monkeypatch.setattr(loaders, "ocr_pdf_page", local_page)
    original = loaders.extract_pdf(_blank_pdf())
    assert original.is_empty and called == []
    result = loaders.extract_pdf(_blank_pdf(), ocr_options=OcrOptions(enabled=True, language="eng"))
    assert result.blocks[0].page_or_sheet == "pagina 1"
    assert "42" in result.plain_text()
    assert called == [(1, "eng")]
    assert any("OCR local" in warning for warning in result.warnings)


def test_missing_ocr_is_actionable_and_never_falls_back_to_remote(monkeypatch):
    monkeypatch.setattr(local_ocr.shutil, "which", lambda name: None)
    result = loaders.extract_pdf(_blank_pdf(), ocr_options=OcrOptions(enabled=True))
    assert result.is_empty
    assert any("pdftoppm y Tesseract" in warning for warning in result.warnings)


def test_isolated_worker_preserves_safe_encrypted_pdf_error():
    with pytest.raises(ExtractionFailedError, match="contrasena"):
        isolated_extraction.extract_document(_blank_pdf(encrypted=True), filename="protegido.pdf")


def test_isolated_worker_extracts_pptx_contract():
    result = isolated_extraction.extract_document(presentation((1,)), filename="diapositiva.pptx")
    assert {block.page_or_sheet for block in result.blocks} == {"diapositiva 1"}
    assert "1,000,000" in result.plain_text()
