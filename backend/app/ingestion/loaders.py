# Creado por Aldo Garcia.
"""Extractores de texto por formato.

Formatos: ``.docx``, ``.pptx``, ``.md``, ``.pdf``, ``.txt``, ``.xlsx``.
Mejora permitida: ``.csv``.

Controles de seguridad aplicados en todos los extractores (seccion 17):

* limites de paginas, hojas, filas y celdas -- un archivo OOXML de 3 KB puede
  expandirse a gigabytes (zip bomb) si no se acota lo que se lee;
* nunca se ejecutan macros, formulas ni contenido incrustado: openpyxl se abre
  con ``data_only=True`` (lee el valor cacheado, no evalua) y python-docx no
  ejecuta nada;
* si una pagina de PDF no tiene texto extraible se **deja constancia** en lugar
  de inventar contenido.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

from app.common.errors import ExtractionFailedError, UnsupportedFileError
from app.common.logging import get_logger
from app.ingestion.local_ocr import OcrOptions, ocr_pdf_page
from app.rag.chunking import TextBlock

logger = get_logger(__name__)

#: Limites duros anti-abuso.
MAX_PDF_PAGES = 500
MAX_XLSX_SHEETS = 50
MAX_XLSX_ROWS_PER_SHEET = 5000
MAX_XLSX_COLUMNS = 100
MAX_CSV_ROWS = 20000
MAX_TEXT_BYTES = 20 * 1024 * 1024
MAX_DOCX_PARAGRAPHS = 20000
MAX_DOCX_TABLES = 1000
MAX_DOCX_TABLE_ROWS = 5000
MAX_DOCX_TABLE_COLUMNS = 100

_EXTENSION_MIME = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".md": "text/markdown",
    ".pdf": "application/pdf",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".txt": "text/plain",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
}


def supported_extensions() -> frozenset[str]:
    return frozenset(_EXTENSION_MIME)


def mime_for_extension(extension: str) -> str:
    return _EXTENSION_MIME.get(extension.lower(), "application/octet-stream")


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    """Resultado de la extraccion: bloques estructurales y avisos."""

    blocks: list[TextBlock]
    mime_type: str
    #: Avisos no fatales (paginas sin texto, hojas truncadas...). Se persisten
    #: para que el operador sepa que parte del documento no se indexo.
    warnings: list[str] = field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        return not any(b.text.strip() for b in self.blocks)

    def plain_text(self) -> str:
        return "\n\n".join(b.text for b in self.blocks)


# ---------------------------------------------------------------------------
# Texto plano y Markdown
# ---------------------------------------------------------------------------
def _decode_text(data: bytes) -> str:
    """Detecta y normaliza el encoding de forma segura.

    Se prueban los encodings habituales en documentos de RH generados en Windows.
    Como ultimo recurso se decodifica con reemplazo para no perder el documento
    entero por un byte invalido.
    """
    # PowerShell/Office pueden guardar TXT con BOM UTF-16. No convertir sus
    # bytes NUL en texto aparentemente valido por el fallback Windows-1252.
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError as exc:
            raise ExtractionFailedError("El archivo UTF-16 esta incompleto o corrupto.") from exc
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_txt(data: bytes) -> ExtractedDocument:
    if len(data) > MAX_TEXT_BYTES:
        raise ExtractionFailedError("El archivo de texto excede el limite de extraccion.")
    from app.rag.chunking import split_into_blocks

    text = _decode_text(data)
    return ExtractedDocument(blocks=split_into_blocks(text), mime_type="text/plain")


def extract_markdown(data: bytes) -> ExtractedDocument:
    """Markdown conserva headings, listas y bloques semanticos."""
    if len(data) > MAX_TEXT_BYTES:
        raise ExtractionFailedError("El archivo Markdown excede el limite de extraccion.")
    from app.rag.chunking import split_into_blocks

    text = _decode_text(data)
    return ExtractedDocument(blocks=split_into_blocks(text), mime_type="text/markdown")


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def extract_docx(data: bytes) -> ExtractedDocument:
    """Extrae DOCX conservando titulos, subtitulos, listas, tablas y orden."""
    try:
        import docx  # python-docx
        from docx.table import Table
    except ImportError as exc:  # pragma: no cover - dependencia declarada
        raise ExtractionFailedError("Falta la dependencia python-docx.", detail=str(exc)) from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionFailedError(
            "El documento Word no pudo abrirse (protegido o corrupto).", detail=str(exc)
        ) from exc

    blocks: list[TextBlock] = []
    warnings: list[str] = []
    section = ""

    # Las colecciones paragraphs/tables separadas pierden la posicion original:
    # una tabla podia heredar el ultimo titulo del documento. python-docx expone
    # los elementos del cuerpo en orden, sin ejecutar contenido incrustado.
    table_index = 0
    paragraph_index = 0
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            table_index += 1
            if table_index > MAX_DOCX_TABLES or len(item.rows) > MAX_DOCX_TABLE_ROWS:
                raise ExtractionFailedError("El documento Word excede el limite de tablas o filas.")
            rows: list[str] = []
            for row in item.rows:
                if len(row.cells) > MAX_DOCX_TABLE_COLUMNS:
                    raise ExtractionFailedError("Una tabla Word excede el limite de columnas.")
                cells = [cell.text.strip().replace("\n", " ").replace("|", "\\|") for cell in row.cells]
                if any(cells):
                    rows.append("| " + " | ".join(cells) + " |")
            if rows:
                blocks.append(
                    TextBlock(
                        "\n".join(rows),
                        section=section or f"Tabla {table_index}",
                        page_or_sheet=f"tabla {table_index}",
                        kind="table",
                    )
                )
            continue
        paragraph = item
        paragraph_index += 1
        if paragraph_index > MAX_DOCX_PARAGRAPHS:
            raise ExtractionFailedError("El documento Word excede el limite de parrafos.")
        locator = f"parrafo {paragraph_index}"
        text = paragraph.text.strip()
        if not text:
            continue
        # ``Paragraph.style`` puede ser ``None`` en un DOCX valido con estilos
        # incompletos; tratarlo como texto normal evita fallar toda la ingesta.
        paragraph_style = paragraph.style
        style = ((paragraph_style.name if paragraph_style is not None else "") or "").lower()
        if style.startswith("heading") or style in ("title", "subtitle"):
            section = text
            # Se emite como heading Markdown para que el chunker lo reconozca
            # como separador estructural de primer nivel.
            level = 1
            if style.startswith("heading"):
                digits = "".join(ch for ch in style if ch.isdigit())
                level = min(6, int(digits)) if digits else 1
            blocks.append(TextBlock(
                f"{'#' * level} {text}", section=section, page_or_sheet=locator, kind="heading", level=level,
            ))
            continue
        kind = "list" if style.startswith("list") else "paragraph"
        prefix = "- " if kind == "list" else ""
        blocks.append(TextBlock(f"{prefix}{text}", section=section, page_or_sheet=locator, kind=kind))

    if not blocks:
        warnings.append("El documento Word no contiene texto extraible.")
    return ExtractedDocument(blocks=blocks, mime_type=_EXTENSION_MIME[".docx"], warnings=warnings)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def extract_pdf(data: bytes, *, ocr_options: OcrOptions | None = None) -> ExtractedDocument:
    """Extrae texto pagina por pagina.

    Una pagina escaneada sin capa de texto produce un aviso explicito. **No** se
    infiere ni se inventa su contenido: Matrix RH prefiere declarar que no puede
    leer una pagina antes que fabricar una politica.
    """
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ExtractionFailedError("Falta la dependencia pypdf.", detail=str(exc)) from exc

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            # Se intenta la contrasena vacia (PDF "protegido" solo contra edicion).
            try:
                if not reader.decrypt(""):
                    raise ExtractionFailedError("El PDF esta protegido con contrasena y no puede procesarse.")
            except Exception as exc:  # noqa: BLE001
                raise ExtractionFailedError(
                    "El PDF esta protegido con contrasena y no puede procesarse.", detail=str(exc)
                ) from exc
    except ExtractionFailedError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionFailedError("El PDF no pudo abrirse.", detail=str(exc)) from exc

    blocks: list[TextBlock] = []
    warnings: list[str] = []
    pages = reader.pages
    if len(pages) > MAX_PDF_PAGES:
        warnings.append(f"PDF truncado a {MAX_PDF_PAGES} paginas de {len(pages)}.")

    from app.rag.chunking import split_into_blocks

    empty_pages: list[int] = []
    options = ocr_options or OcrOptions()
    for page_number, page in enumerate(pages[:MAX_PDF_PAGES], start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"Pagina {page_number}: error de extraccion ({type(exc).__name__}).")
            continue
        if not text.strip():
            if options.enabled:
                try:
                    text = ocr_pdf_page(data, page_number, options)
                    warnings.append(
                        f"Pagina {page_number}: texto obtenido por OCR local; "
                        "revisar cifras y tablas contra el original."
                    )
                except ExtractionFailedError as exc:
                    warnings.append(f"Pagina {page_number}: {exc.message}")
            if not text.strip():
                empty_pages.append(page_number)
                continue
        label = f"pagina {page_number}"
        for block in split_into_blocks(text, default_section=label):
            blocks.append(
                TextBlock(
                    text=block.text,
                    section=block.section or label,
                    page_or_sheet=label,
                    kind=block.kind,
                )
            )

    if empty_pages:
        warnings.append(
            "Paginas sin texto extraible (posible escaneo sin OCR): "
            + ", ".join(str(p) for p in empty_pages[:25])
        )
        if not options.enabled:
            warnings.append("El OCR local esta desactivado; las paginas escaneadas no se indexaron.")
    return ExtractedDocument(blocks=blocks, mime_type=_EXTENSION_MIME[".pdf"], warnings=warnings)


def extract_pptx(data: bytes) -> ExtractedDocument:
    from app.ingestion.pptx import extract_pptx_blocks

    blocks, warnings = extract_pptx_blocks(data)
    return ExtractedDocument(blocks, _EXTENSION_MIME[".pptx"], warnings)


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------
def extract_xlsx(data: bytes) -> ExtractedDocument:
    """Procesa el workbook hoja por hoja.

    Estrategia: se conserva el nombre de la hoja y los encabezados, y cada fila se
    serializa como ``columna: valor``. Es lo que permite que una consulta
    semantica sobre "dias de vacaciones por antiguedad" recupere la fila correcta
    en lugar de una tabla entera sin contexto.
    """
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]
    except ImportError as exc:  # pragma: no cover
        raise ExtractionFailedError("Falta la dependencia openpyxl.", detail=str(exc)) from exc

    try:
        # data_only=True: se lee el valor cacheado. NUNCA se evaluan formulas.
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionFailedError(
            "El libro de Excel no pudo abrirse (protegido o corrupto).", detail=str(exc)
        ) from exc

    blocks: list[TextBlock] = []
    warnings: list[str] = []
    try:
        sheet_names = workbook.sheetnames[:MAX_XLSX_SHEETS]
        if len(workbook.sheetnames) > MAX_XLSX_SHEETS:
            warnings.append(f"Libro truncado a {MAX_XLSX_SHEETS} hojas.")

        for sheet_name in sheet_names:
            sheet = workbook[sheet_name]
            blocks.append(
                TextBlock(f"## Hoja: {sheet_name}", section=sheet_name, page_or_sheet=sheet_name, kind="heading")
            )
            headers: list[str] = []
            row_lines: list[str] = []
            if sheet.max_column and sheet.max_column > MAX_XLSX_COLUMNS:
                warnings.append(f"Hoja '{sheet_name}' truncada a {MAX_XLSX_COLUMNS} columnas.")
            read_rows = min(sheet.max_row or MAX_XLSX_ROWS_PER_SHEET + 1, MAX_XLSX_ROWS_PER_SHEET + 1)
            read_columns = min(sheet.max_column or MAX_XLSX_COLUMNS, MAX_XLSX_COLUMNS)

            # Limitar dentro del lector: cortar row[:N] despues de iter_rows
            # ya materializaba todas las columnas de una fila. Contar filas
            # fisicas evita recorrer millones de huecos en libros dispersos.
            # Una fila adicional permite informar el truncamiento al operador.
            for row_number, row in enumerate(sheet.iter_rows(
                values_only=True, max_row=read_rows, max_col=read_columns,
            ), start=1):
                if row_number > MAX_XLSX_ROWS_PER_SHEET:
                    warnings.append(
                        f"Hoja '{sheet_name}' truncada a {MAX_XLSX_ROWS_PER_SHEET} filas."
                    )
                    break
                values = ["" if v is None else str(v).strip() for v in row[:MAX_XLSX_COLUMNS]]
                if not any(values):
                    continue
                if not headers:
                    headers = [v or f"col{i + 1}" for i, v in enumerate(values)]
                    continue
                pairs = [
                    f"{headers[i]}: {value}"
                    for i, value in enumerate(values)
                    if i < len(headers) and value
                ]
                if pairs:
                    row_lines.append("- " + "; ".join(pairs))

            if headers:
                blocks.append(
                    TextBlock(
                        "Columnas: " + ", ".join(headers),
                        section=sheet_name,
                        page_or_sheet=sheet_name,
                        kind="paragraph",
                    )
                )
            if row_lines:
                blocks.append(
                    TextBlock(
                        "\n".join(row_lines),
                        section=sheet_name,
                        page_or_sheet=sheet_name,
                        kind="list",
                    )
                )
    finally:
        workbook.close()

    if not blocks:
        warnings.append("El libro de Excel no contiene datos extraibles.")
    return ExtractedDocument(blocks=blocks, mime_type=_EXTENSION_MIME[".xlsx"], warnings=warnings)


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------
def extract_csv(data: bytes) -> ExtractedDocument:
    """CSV con deteccion de delimitador y el mismo formato fila -> ``columna: valor``."""
    if len(data) > MAX_TEXT_BYTES:
        raise ExtractionFailedError("El archivo CSV excede el limite de extraccion.")
    text = _decode_text(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel  # type: ignore[assignment]

    reader = csv.reader(io.StringIO(text), dialect)
    blocks: list[TextBlock] = []
    warnings: list[str] = []
    headers: list[str] = []
    lines: list[str] = []

    for index, row in enumerate(reader):
        if index >= MAX_CSV_ROWS:
            warnings.append(f"CSV truncado a {MAX_CSV_ROWS} filas.")
            break
        values = [cell.strip() for cell in row]
        if not any(values):
            continue
        if not headers:
            headers = [v or f"col{i + 1}" for i, v in enumerate(values)]
            continue
        pairs = [f"{headers[i]}: {v}" for i, v in enumerate(values) if i < len(headers) and v]
        if pairs:
            lines.append("- " + "; ".join(pairs))

    if headers:
        blocks.append(TextBlock("Columnas: " + ", ".join(headers), section="csv", kind="paragraph"))
    if lines:
        blocks.append(TextBlock("\n".join(lines), section="csv", kind="list"))
    return ExtractedDocument(blocks=blocks, mime_type="text/csv", warnings=warnings)


# ---------------------------------------------------------------------------
# Despachador
# ---------------------------------------------------------------------------
_EXTRACTORS = {
    ".txt": extract_txt,
    ".md": extract_markdown,
    ".docx": extract_docx,
    ".pdf": extract_pdf,
    ".pptx": extract_pptx,
    ".xlsx": extract_xlsx,
    ".csv": extract_csv,
}


def extract_document(data: bytes, *, filename: str, ocr_options: OcrOptions | None = None) -> ExtractedDocument:
    """Selecciona el extractor por extension. Extension desconocida = rechazo."""
    extension = Path(filename).suffix.lower()
    if extension in {".doc", ".ppt"}:
        raise UnsupportedFileError(
            "El formato Office antiguo requiere conversion local a DOCX o PPTX; no hay conversor automatico habilitado."
        )
    extractor = _EXTRACTORS.get(extension)
    if extractor is None:
        raise UnsupportedFileError(f"Extension no soportada: {extension or '(sin extension)'}")
    result = extract_pdf(data, ocr_options=ocr_options) if extension == ".pdf" else extractor(data)
    if result.warnings:
        logger.info(
            "ingestion.extraction_warnings",
            extra={"extraction_warnings": result.warnings[:5], "file_extension": extension},
        )
    return result
